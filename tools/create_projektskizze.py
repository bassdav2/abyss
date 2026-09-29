#!/usr/bin/env python3
"""Erzeugt die Projektskizze (PM3, Meilenstein M1) aus einer einzigen Inhaltsquelle.

Ausgaben in docs/m1-projektskizze:
  ABYSS_Projektskizze_M1.pdf   Layout nach SoE-Dokumentvorlage, gerendert mit Chrome (headless)
  ABYSS_Projektskizze_M1.docx  auf Basis der offiziellen SoE-Vorlage, zum Weiterbearbeiten in Word
  PROJEKTSKIZZE.md             Lesefassung für GitHub
  abbildungen/*.png            Abbildungen aus echten Spielszenen und Diagrammen

Voraussetzungen: Pillow, python-docx, Google Chrome, pdftotext (poppler), rsvg-convert, Spielszenen aus
  ./gradlew sceneShot --args="build/scenes 0,4,9,14,19 4"

Aufruf aus dem Projektordner:  python3 tools/create_projektskizze.py
Teamangaben (Nummer, Namen) stehen in TEAM und MEMBERS unten.
"""
from __future__ import annotations

import html
import re
import shutil
import subprocess
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'm1-projektskizze'
FIG = OUT / 'abbildungen'
BUILD = ROOT / 'build' / 'skizze'
SCENES = ROOT / 'build' / 'scenes'
TEMPLATE = ROOT.parent / 'Kursunterlagen' / 'Software-Projekt 3 HS26_20260914_0829' / 'Allgemeines' / 'Sprachlicher Teil' / '01_Dokumentvorlage_SoE_Projektskizze.docx'
CHROME = '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'
WORD_FONTS = Path('/Applications/Microsoft Word.app/Contents/Resources/DFonts')
STEM = 'ABYSS_Projektskizze_M1'

# ------------------------------------------------------------------ Teamangaben (vom Team zu ergänzen)
TEAM = 'Team ABYSS'
MEMBERS = ['David Bass']
MEMBERS_TODO = 'weitere Teammitglieder in alphabetischer Reihenfolge ergänzen'
DATE = '29. September 2026'
PRODUCT = 'ABYSS – Vom Heck bis zur Brücke'
SUBTITLE = 'Ein 2D-Action-Roguelite in Java und JavaFX'
ACCESSED = 'Sep. 26, 2026'

# ------------------------------------------------------------------ Quellen (IEEE, Nummer nach erster Nennung)
REFS = {
    'harris': 'J. Harris, <i>Exploring Roguelike Games</i>. Boca Raton, FL, USA: CRC Press, 2020.',
    'deadcells_sales': f'C. Kerr, “Dead Cells has topped 10 million sales worldwide,” <i>Game Developer</i>, Jun. 5, 2023. [Online]. Available: https://www.gamedeveloper.com/business/dead-cells-has-topped-10-million-sales-worldwide. [Accessed: {ACCESSED}].',
    'hades_sales': f'Supergiant Games, “Hades has now sold more than 1,000,000 copies,” X (formerly Twitter), Sep. 20, 2020. [Online]. Available: https://x.com/SupergiantGames/status/1307744738552938496. [Accessed: {ACCESSED}].',
    'ftl': f'Subset Games, <i>FTL: Faster Than Light</i>. [Computer game]. Subset Games, 2012. [Online]. Available: https://store.steampowered.com/app/212680/. [Accessed: {ACCESSED}].',
    'barotrauma': f'FakeFish and Undertow Games, <i>Barotrauma</i>. [Computer game]. Daedalic Entertainment, 2023. [Online]. Available: https://barotraumagame.com/. [Accessed: {ACCESSED}].',
    'snowpiercer': 'J. Bong, <i>Snowpiercer</i>. [Film]. CJ Entertainment, 2013.',
    'deadcells': 'Motion Twin, <i>Dead Cells</i>. [Computer game]. Motion Twin, 2018.',
    'hades': 'Supergiant Games, <i>Hades</i>. [Computer game]. Supergiant Games, 2020.',
    'libgdx': f'libGDX, “libGDX – a cross-platform Java game development framework based on OpenGL (ES),” libgdx.com. [Online]. Available: https://libgdx.com/. [Accessed: {ACCESSED}].',
    'openjfx': f'OpenJFX, “JavaFX – an open source, next generation client application platform,” openjfx.io. [Online]. Available: https://openjfx.io/. [Accessed: {ACCESSED}].',
    'nystrom': f'R. Nystrom, <i>Game Programming Patterns</i>. Genever Benning, 2014. [Online]. Available: https://gameprogrammingpatterns.com/. [Accessed: {ACCESSED}].',
    'kickoff': 'ZHAW School of Engineering, “Einführung und Kick-off,” Folien Software-Projekt 3 (PM3), Ausgabe HS26, Winterthur, 2026.',
    'pm': 'ZHAW School of Engineering, “Projektmanagement,” Folien zum Kick-off Software-Projekt 3 (PM3), Winterthur, 2026.',
    'ki_pm3': 'ZHAW School of Engineering, “Rahmenbedingungen zum Einsatz von KI im PM3,” Version 1.0, Kursunterlagen HS26, Winterthur, 2026.',
    'ki_soe': 'ZHAW School of Engineering, “Anhang zur ZHAW-Richtlinie «KI bei Leistungsnachweisen»: Deklarationspflicht von generativer KI bei Arbeiten an der SoE,” Version 1.0.0, Winterthur, Feb. 2024.',
}

# ------------------------------------------------------------------ Inhalt
# Blöcke: ('h1'|'h2', Titel) · ('p', Text) · ('ul', [Punkte]) · ('fig', Schlüssel, Datei, Legende, Breite in cm)
#         ('tab', Schlüssel, Legende, Kopfzeile, Zeilen, Spaltenanteile) · ('persona', Text)
# Im Text: **fett**, *kursiv*, {ref:x} → [n], {fig:x} → Abbildung n, {tab:x} → Tabelle n

CONTENT = [
    ('h1', 'Ausgangslage'),
    ('p', 'Kurze, in sich abgeschlossene Spielsitzungen sind ein verbreitetes Muster bei Computerspielen für Desktop-Rechner. Ein Genre, das darauf aufbaut, sind sogenannte **Roguelites**: Jeder Durchlauf wird zufällig zusammengestellt, endet mit Sieg oder Niederlage und beginnt danach von vorn, wobei ein Teil des Fortschritts über die Versuche hinweg erhalten bleibt {ref:harris}. Wie gross die Nachfrage nach solchen Spielen ist, zeigen zwei Beispiele: *Dead Cells* wurde nach Angaben der Entwickler mehr als zehn Millionen Mal verkauft {ref:deadcells_sales}, *Hades* überschritt kurz nach dem Erscheinen der Vollversion eine Million verkaufte Exemplare {ref:hades_sales}.'),
    ('p', 'Eine zweite Gruppe von Spielen nutzt ein einzelnes Fahrzeug als Spielwelt. In *FTL: Faster Than Light* führen Spielende ein Raumschiff durch eine zufällig erzeugte Galaxie {ref:ftl}, in *Barotrauma* bedient eine Mannschaft gemeinsam ein U-Boot {ref:barotrauma}. Im Mittelpunkt steht dort die Steuerung des Fahrzeugs und seiner Systeme, nicht der Weg einer einzelnen Figur durch das Innere des Fahrzeugs.'),
    ('p', 'Das Projektteam entwickelt im Modul Software-Projekt 3 (PM3) innerhalb eines Semesters eine Anwendung in Java mit selbst definierter, mindestens geschichteter Architektur; Frameworks und relationale Datenbanken sind ausgeschlossen {ref:kickoff}. Einen externen Auftraggeber gibt es nicht. Die Zielgruppe wird deshalb über eine Persona beschrieben, deren Annahmen im Projektverlauf mit Testpersonen überprüft werden (vgl. Kapitel 3).'),

    ('h1', 'Idee'),
    ('p', 'ABYSS ist ein zweidimensionales Action-Roguelite, das vollständig im Inneren eines sehr langen U-Boots spielt. Die Spielfigur gehört zur Besatzung und beginnt im Heck. Raum für Raum kämpft sie sich bis zur Brücke vor, um die Kontrolle über das Boot zu übernehmen. Die räumliche Ordnung – hinten enge Wartungsräume, vorne die Schiffsführung – ist vom Film *Snowpiercer* inspiriert {ref:snowpiercer}; Figuren, Namen und Handlung sind eigenständig.'),
    ('p', 'Wie {fig:boot} zeigt, ist das Boot in vier Sektionen mit je sechs Räumen gegliedert. Jede Sektion endet mit einem Wächter, einem besonders starken Gegner. Vor jedem Raum wählen Spielende an einer Abzweigung zwischen zwei Wegen, etwa zwischen einem Kampfraum mit höherer Belohnung und einem ruhigeren Versorgungsraum. Ein Durchlauf, im Spiel «Tauchgang» genannt, ist auf eine Sitzung von etwa 30 Minuten ausgelegt; diese Dauer ist eine Annahme und wird in Spieltests überprüft.'),
    ('p', 'Die Idee geht von folgendem Problem aus: Spielende mit wenig Zeit suchen Spiele, die innerhalb einer Sitzung einen vollständigen Spannungsbogen bieten und deren Niederlagen nachvollziehbar sind. Der Kernnutzen von ABYSS ist deshalb ein abgeschlossener, fairer Durchlauf pro Sitzung, der sich bei jedem Versuch anders spielt, weil Route, Räume, Gegner und Belohnungen neu zusammengestellt werden.'),
    ('fig', 'boot', 'abb-boot.png', 'Aufbau des U-Boots: 24 Räume in vier Sektionen vom Heck bis zur Brücke; jede Sektion endet mit einem Wächter (Bootskarte im Prototyp, eigene Beschriftung)', 16),
    ('p', 'Eine Bildschirmansicht des bestehenden Prototyps zeigt {fig:screen}. Die nummerierten Markierungen bezeichnen (1) Integrität, Energie und Reparatursets der Spielfigur, (2) die aktuelle Angriffswelle, (3) die Raumposition 15 von 24 sowie gesammelten Schrott und Datenkerne, (4) die Spielfigur mit Stirnlampe, (5) die rote Markierung, mit der ein Gegner seinen Angriff ankündigt, (6) Waffe, aktives Modul und Ausweichen sowie (7) die installierten Module. Alle Grafiken werden im Programmcode als Pixelgrafik mit 480 × 270 Bildpunkten erzeugt und für die Anzeige ganzzahlig vergrössert.'),
    ('fig', 'screen', 'abb-bildschirm.png', 'Bildschirmansicht des Prototyps im Forschungsdeck mit nummerierten Bedienelementen', 16),

    ('h1', 'Kundennutzen'),
    ('p', 'Die Hauptnutzenden sind Einzelspielende, die am eigenen Laptop oder Desktop-Rechner in kurzen Pausen spielen. Stellvertretend für diese Gruppe steht die folgende Persona. Sie ist eine Annahme des Teams; ihre Merkmale werden in Interviews und Spieltests überprüft.'),
    ('persona', '**Persona Lina (22)** studiert und spielt am Laptop zwischen Vorlesungen oder am Abend, meist 20 bis 30 Minuten am Stück. Sie mag kurze Actionspiele und abwechslungsreiche Ausrüstung. Niederlagen akzeptiert sie, wenn sie deren Ursache erkennt. Lange Einführungen und einen Zwang zur Internetverbindung lehnt sie ab.'),
    ('p', 'Für Lina und vergleichbare Spielende ergibt sich folgender Nutzen:'),
    ('ul', [
        '**Abgeschlossene Sitzung:** Ein Durchlauf passt in eine Pause. Beim Beenden wird der Eingang des aktuellen Raums gesichert, sodass eine Unterbrechung ohne Verlust möglich ist.',
        '**Nachvollziehbare Niederlagen:** Gegner kündigen ihre Angriffe sichtbar an (Markierung 5 in {fig:screen}). Eine Niederlage lässt sich dadurch auf eine konkrete Situation zurückführen.',
        '**Abwechslung statt reiner Zahlensteigerung:** Route, Räume, Raumzustände und Belohnungen werden pro Durchlauf neu zusammengestellt. Dauerhafte Freischaltungen erweitern die Auswahl an Figuren, Waffen und Modulen, statt nur Werte zu erhöhen.',
        '**Eigene Entscheidungen:** Die Routenwahl verlangt eine Abwägung zwischen Risiko und Belohnung. Anlagen im Raum, etwa Hydraulikpressen, lassen sich gezielt gegen Gegner einsetzen.',
        '**Einstieg ohne Hürden:** Das Spiel läuft lokal ohne Konto und ohne Internetverbindung. Der erste Raum führt ohne Texttafeln in die Steuerung ein.',
    ]),

    ('h1', 'Stand der Technik / Konkurrenzanalyse'),
    ('p', 'Die in {tab:konkurrenz} aufgeführten Spiele decken jeweils einen Teil der Idee ab. Ausgewählt wurden zwei kommerziell erfolgreiche Roguelites und zwei Spiele, deren Welt aus einem einzelnen Fahrzeug besteht.'),
    ('tab', 'konkurrenz', 'Einordnung von ABYSS gegenüber verwandten Spielen (eigene Darstellung)',
     ['Spiel', 'Entwicklung, Jahr', 'Spielprinzip', 'Gemeinsamkeit mit ABYSS', 'Unterschied zu ABYSS'],
     [['Dead Cells {ref:deadcells}', 'Motion Twin, 2018', '2D-Action-Plattformer mit zufälligen Durchläufen', 'Plattformkampf, Neubeginn nach Niederlage, dauerhafte Freischaltungen', 'Welt ist ein verzweigtes Schloss, keine lineare Reise durch ein Fahrzeug'],
      ['Hades {ref:hades}', 'Supergiant Games, 2020', 'Action-Roguelite in Draufsicht', 'Raumfolge mit Belohnungswahl nach jedem Raum', 'Draufsicht statt Seitenansicht, mythologische Unterwelt statt Maschinenwelt'],
      ['FTL: Faster Than Light {ref:ftl}', 'Subset Games, 2012', 'Weltraum-Strategiespiel mit Roguelike-Elementen', 'Routenwahl, Entscheidungen unter Risiko, ein Fahrzeug als Welt', 'Steuerung des Schiffs statt Kampf einer Figur'],
      ['Barotrauma {ref:barotrauma}', 'FakeFish und Undertow Games, 2023', 'Kooperative U-Boot-Simulation', 'U-Boot als Spielwelt in 2D-Seitenansicht', 'Mehrspieler-Simulation ohne Roguelite-Durchläufe']],
     [.17, .15, .2, .24, .24]),
    ('p', 'Aus {tab:konkurrenz} geht hervor, dass keines der Vergleichsspiele den Plattformkampf einer einzelnen Figur mit einer linearen Reise durch das Innere eines Fahrzeugs verbindet. ABYSS setzt an dieser Stelle an: Das U-Boot ist zugleich Spielwelt und Fortschrittsanzeige, und seine Technik – Förderbänder, Pressen, Lasergitter und Notschalter – wird Teil des Kampfes. Als Studienprojekt tritt ABYSS nicht in kommerzielle Konkurrenz zu diesen Spielen; die Einordnung dient der Abgrenzung des Umfangs.'),
    ('p', 'Technisch werden Spiele in Java häufig mit einem Spiel-Framework wie libGDX umgesetzt, das Grafik, Eingabe und Audio auf Basis von OpenGL für mehrere Plattformen bereitstellt {ref:libgdx}. Ein solches Framework kommt für ABYSS nicht in Frage, weil die Randbedingungen des Moduls Frameworks ausschliessen und JavaFX als erste Wahl für die Oberfläche nennen {ref:kickoff}. ABYSS verwendet deshalb die Client-Plattform JavaFX {ref:openjfx} und eine eigene Spielschicht. Spielregeln, Darstellung und Speicherung werden als getrennte Schichten selbst entworfen, getestet und im Technischen Bericht begründet. Bewährte Entwurfsmuster für Spielschleife und Zustandsautomaten sind in der Literatur beschrieben {ref:nystrom}.'),

    ('h1', 'Hauptablauf (Kontextszenario)'),
    ('p', 'Lina hat nach einer Vorlesung eine halbe Stunde Zeit. Sie startet das Spiel auf ihrem Laptop und beginnt einen neuen Tauchgang. In der Schleuse wählt sie eine von mehreren Figuren; eine kurze Beschreibung nennt jeweils deren Stärken und Schwächen. Sie entscheidet sich für die Mechanikerin, weil deren Werte ausgewogen sind.'),
    ('p', 'Der erste Raum im Heck ist ruhig und zeigt, wie sich die Figur bewegt, springt und ausweicht. Im zweiten Raum trifft Lina auf eine Patrouille. Ein Gegner leuchtet kurz rot auf, bevor er zuschlägt; Lina weicht aus und besiegt ihn. Sobald alle Gegner besiegt sind, öffnet sich das Schott zum nächsten Raum, und in einer Bergungskapsel werden drei Module angeboten. Lina wählt eines, das ihren Angriff verstärkt.'),
    ('p', 'An der nächsten Abzweigung zeigen zwei Bildschirme, was hinter den Schotts liegt: ein Frachtraum mit zwei Gegnerwellen und ein Depot mit Vorräten. Da die Integrität ihrer Figur gesunken ist, wählt sie das Depot und repariert ihre Ausrüstung. Einige Räume später arbeitet in einem Maschinenraum eine Hydraulikpresse. Lina lockt einen gepanzerten Gegner unter die Presse und nutzt die Anlage so zu ihrem Vorteil.'),
    ('p', 'Am Ende der Hecksektion tritt sie gegen den ersten Wächter an und verliert knapp, weil sie dessen Stampfangriff zu spät erkennt. Die Auswertung zeigt, wie weit sie gekommen ist und wie viele Datenkerne sie gesammelt hat. Im Archiv schaltet sie damit eine weitere Figur frei. Nach knapp dreissig Minuten beendet Lina das Spiel mit dem Vorsatz, beim nächsten Versuch den Angriff früher zu erkennen und eine andere Route zu wählen.'),
    ('p', '{fig:ablauf} fasst diesen Ablauf schematisch zusammen. Der Kern ist die Schleife aus Raum, Bergung und Routenwahl, die sich bis zur Brücke wiederholt. Sieg und Niederlage führen beide zur Auswertung und ins Archiv, von wo aus der nächste Tauchgang beginnt.'),
    ('fig', 'ablauf', 'abb-ablauf.png', 'Ablauf eines Tauchgangs mit der Schleife aus Raum, Bergung und Routenwahl (eigene Darstellung)', 16),

    ('h1', 'Weitere Anforderungen'),
    ('h2', 'Funktionale Anforderungen'),
    ('p', 'Neben dem Hauptablauf sind folgende Funktionen vorgesehen:'),
    ('ul', [
        'Plattformbewegung mit Laufen, Sprung, Luftsprung, Ausweichen und Fallen durch Laufstege',
        'Nahkampf mit Kombinationen, Fernkampfwaffen und aktive Module, die Energie verbrauchen',
        'vier Sektionen mit je sechs Räumen, Routenwahl mit Vorschau und Sonderräume (Händler, Werkstatt, Kapelle, Versorgung)',
        'Gegner mit angekündigten Angriffen, Elite-Varianten und ein Wächter pro Sektion mit mehreren Phasen',
        'Raumzustände wie Stromausfall, Hüllenbruch oder Schlagseite, die in der Routenwahl angekündigt werden',
        'Raumtechnik (Förderbänder, Dampfdüsen, Pressen, Lasergitter, Notschalter), die Figur und Gegner betrifft',
        'Module in mehreren Seltenheitsstufen und Zusatzboni für bestimmte Modulpaare («Resonanzen»)',
        'dauerhafter Fortschritt über ein Archiv mit Freischaltungen, eine Garderobe und ein Logbuch mit Zielen',
        'Sicherung am Raumeingang, Fortsetzen nach dem Beenden und ein lokales Profil',
        'Optionen für Lautstärke, Bildschirmmodus und Effekte; Bedienung mit Maus und Tastatur',
    ]),
    ('h2', 'Nicht-funktionale Anforderungen'),
    ('p', '{tab:nfr} fasst die nicht-funktionalen Anforderungen zusammen. Zu jeder Anforderung ist ein Kriterium angegeben, mit dem sich am Projektende prüfen lässt, ob sie erfüllt ist.'),
    ('tab', 'nfr', 'Nicht-funktionale Anforderungen mit Prüfkriterium',
     ['ID', 'Anforderung', 'Prüfkriterium'],
     [['Q-01', 'Testbarkeit: Die Spielregeln sind unabhängig von Oberfläche und Dateisystem.', 'Automatische Tests der Spiellogik laufen ohne Fenster.'],
      ['Q-02', 'Robustheit: Beschädigte oder veraltete Spielstände verhindern den Start nicht.', 'Mit einer defekten Datei startet das Spiel mit neuem Profil; das Original bleibt als Sicherung erhalten.'],
      ['Q-03', 'Performance: Die Darstellung läuft flüssig.', 'Die Berechnung eines Bildes dauert im Mittel unter 16 ms (60 Bilder pro Sekunde) auf einem aktuellen Laptop.'],
      ['Q-04', 'Bedienbarkeit: Angriffe sind erkennbar, der Einstieg gelingt ohne Anleitung.', 'Vier von fünf Testpersonen erkennen eine Angriffsankündigung ohne Erklärung und schliessen den ersten Raum ohne Hilfe ab.'],
      ['Q-05', 'Nachvollziehbarkeit: KI-Einsatz, Quellen und Tests sind dokumentiert.', 'Deklaration nach SoE-Richtlinie, Quellen nach IEEE, Testergebnisse im Repository.'],
      ['Q-06', 'Portabilität: Das Spiel läuft auf den Rechnern des Teams.', 'Start unter macOS, Windows und Linux auf allen Teamrechnern.'],
      ['Q-07', 'Datenschutz: Es werden keine Daten übertragen.', 'Kein Netzwerkzugriff; Spielstände liegen ausschliesslich lokal.']],
     [.09, .45, .46]),

    ('h1', 'Weiterführende Ideen'),
    ('p', 'Über den Semesterumfang hinaus kann sich die Anwendung in folgende Richtungen entwickeln:'),
    ('ul', [
        'Unterstützung von Gamepads, wofür eine zusätzliche Bibliothek nötig ist, deren Einsatz mit den Dozierenden abzusprechen wäre',
        'frei belegbare Tasten und weitere Optionen zur Barrierefreiheit, etwa Farbschemata für Farbfehlsichtige',
        'weitere Sektionen, Gegner und Wächter sowie Story-Fragmente mit alternativen Enden',
        'eine Bestenliste für den täglichen Tauchgang mit gemeinsamem Startwert, die eine Serverkomponente erfordert',
        'eine Veröffentlichung auf einer Vertriebsplattform nach Abschluss des Semesters',
    ]),

    ('h1', 'Ressourcen'),
    ('p', '**Team und Fähigkeiten.** Das Team besteht aus Studierenden im dritten Semester des Bachelorstudiengangs Informatik. Alle Mitglieder bringen Java-Kenntnisse aus den vorangehenden Programmiermodulen sowie Grundlagen der Analyse und des Entwurfs aus dem parallel besuchten Modul Software-Entwicklung 1 mit. Die Rollen werden entlang der Pakete der Anwendung verteilt: Spiellogik, Darstellung, Oberfläche, Speicherung sowie Test und Qualität.'),
    ('p', '**Fehlendes Know-how.** Spielphysik und Kollisionserkennung, Echtzeitdarstellung, Spieldesign mit Balancing sowie die Durchführung von Nutzertests sind für das Team neu. Diese Themen werden früh im Prototyp erprobt und mit Literatur zu Entwurfsmustern für Spiele {ref:nystrom} erarbeitet.'),
    ('p', '**Bestehender Prototyp.** Ein technischer Prototyp (Version 1.1) liegt bereits vor. Er wurde mit Unterstützung generativer KI erstellt (vgl. Anhang) und belegt, dass die Kernmechaniken in Java und JavaFX umsetzbar sind. Im Semester erarbeitet sich das Team diesen Code, prüft und testet ihn und entwickelt ihn gezielt weiter. Dieser Einarbeitungsaufwand ist in der Schätzung berücksichtigt.'),
    ('p', '**Aufwand.** Gemäss Kursvorgabe stehen pro Person rund 120 Stunden bis zur Semesterwoche 13 zur Verfügung {ref:kickoff}. Bei angenommen fünf Teammitgliedern ergibt sich ein Rahmen von etwa 600 Stunden, der sich wie in {tab:aufwand} verteilt.'),
    ('tab', 'aufwand', 'Grobe Aufwandsschätzung für das Semester bei fünf Teammitgliedern',
     ['Arbeitspaket', 'Anteil', 'Stunden'],
     [['Projektmanagement, Reviews und Präsentationen', '15 %', '90'],
      ['Analyse und Anforderungen (Use Cases, Domänenmodell, Interviews)', '10 %', '60'],
      ['Einarbeitung in den Prototyp, Architektur und Design', '20 %', '120'],
      ['Implementation und Erweiterung', '30 %', '180'],
      ['Test, Spieltests und Balancing', '10 %', '60'],
      ['Technische Berichte und Dokumentation', '15 %', '90'],
      ['**Total**', '**100 %**', '**600**']],
     [.7, .15, .15]),
    ('p', 'Innerhalb der 14 Semesterwochen sind damit ein vollständiger, getesteter Durchlauf durch alle vier Sektionen, die Berichte und Nutzertests mit Personen aus dem Umfeld realistisch. Die weiterführenden Ideen aus Kapitel 7 sowie Tests auf weiteren Plattformen folgen erst nach dem Semester. An Sachmitteln werden nur die eigenen Rechner und kostenlose Werkzeuge (Git, Gradle, eine Java-Entwicklungsumgebung) benötigt. Die UI-Technologie JavaFX und allfällige Bibliotheken werden gemäss den Randbedingungen mit der Fachdozentin oder dem Fachdozenten abgesprochen.'),

    ('h1', 'Risiken'),
    ('p', '{tab:risiken} nennt die Risiken, die über das übliche Mass eines Semesterprojekts hinausgehen, mit Einschätzung und geplanter Massnahme.'),
    ('tab', 'risiken', 'Risiken mit Eintrittswahrscheinlichkeit, Auswirkung und Massnahme',
     ['Nr.', 'Risiko', 'Wahrsch.', 'Auswirk.', 'Massnahme'],
     [['R1', 'Das Team versteht den KI-gestützt erstellten Prototyp nicht ausreichend und kann Entscheide nicht begründen.', 'hoch', 'hoch', 'Verantwortung pro Paket, Code-Reviews; jede Person erklärt ihr Paket im Review; Änderungen nur mit eigenem Test'],
      ['R2', 'Der Umfang wächst über das Semester hinaus.', 'mittel', 'hoch', 'Use Cases priorisieren; neue Inhalte erst nach Erreichen der Iterationsziele'],
      ['R3', 'Steuerung und Fairness überzeugen die Testpersonen nicht.', 'mittel', 'hoch', 'ab Iteration 3 Spieltests mit drei bis fünf Personen, Anpassung nach Beobachtung'],
      ['R4', 'Die Akzeptanz des KI-Einsatzes durch die Dozierenden ist unklar.', 'offen', 'hoch', 'frühe Abstimmung, vollständige Deklaration nach SoE-Richtlinie'],
      ['R5', 'Die Echtzeitdarstellung ist auf schwächeren Rechnern zu langsam.', 'niedrig', 'mittel', 'geringe interne Auflösung, Messung der Zeichenzeit in jedem Build'],
      ['R6', 'Unterschiedliche Betriebssysteme im Team führen zu Problemen.', 'mittel', 'mittel', 'plattformneutraler Java-Code, Build mit Gradle, Test auf allen Teamrechnern'],
      ['R7', 'Teammitglieder fallen aus oder sind unterschiedlich verfügbar.', 'mittel', 'mittel', 'Wissen verteilen (Pair Programming), Aufgaben im Iterationsplan sichtbar machen']],
     [.06, .36, .1, .1, .38]),
    ('p', 'Die Risiken R1 und R4 hängen mit dem KI-gestützten Prototyp zusammen. Sie werden deshalb zuerst bearbeitet: Die Abstimmung erfolgt unmittelbar nach M1, die Einarbeitung ist das Hauptziel der Iteration 2 (vgl. {tab:plan}).'),

    ('h1', 'Grobplanung'),
    ('p', 'Das Projekt folgt dem Kursprozess mit sechs Iterationen zu je zwei Wochen und drei Meilensteinen {ref:pm}. {tab:plan} zeigt für jede Iteration das Ziel, die bearbeiteten Use Cases und den Abschluss. Die Use Cases sind in Kapitel 10.1 beschrieben.'),
    ('tab', 'plan', 'Grobplanung mit Iterationen, Zielen und Meilensteinen (SW = Semesterwoche)',
     ['Iteration', 'Zeitraum', 'Ziel', 'Use Cases', 'Abschluss'],
     [['1', 'SW 2–3', 'Idee, Anforderungen und Projektskizze; Prototyp läuft auf allen Teamrechnern', 'Übersicht', 'M1 Projektskizze, SW 3 (ab 28.09.2026)'],
      ['2', 'SW 4–5', 'Kernablauf verstehen und absichern: Bewegung, Kampf, erster Raum; Code-Walkthrough je Paket', 'UC-01, UC-02', 'Review, SW 5'],
      ['3', 'SW 6–7', 'Durchlauf mit Routenwahl, Bergung und Sicherung; erste Spieltests', 'UC-02, UC-03, UC-05', 'Review, SW 7'],
      ['4', 'SW 8–9', 'Architektur belegen: Domänenmodell, Design, Systemoperationen, Technischer Bericht I', 'UC-04, UC-06, UC-07', 'M2 Lösungsarchitektur, SW 9 (ab 09.11.2026)'],
      ['5', 'SW 10–11', 'Inhalte und dauerhafter Fortschritt, Balancing nach Spieltests', 'UC-08, UC-09, UC-10', 'Review, SW 11'],
      ['6', 'SW 12–13', 'Stabilisierung, Optionen, Technischer Bericht II, Abschlusspräsentation', 'UC-11', 'M3 Prototyp, SW 13 (ab 07.12.2026)']],
     [.1, .12, .4, .16, .22]),
    ('h2', 'Use Cases'),
    ('p', '{fig:uc} zeigt die Use Cases im Überblick. Alle Anwendungsfälle werden von einer einzigen Akteurin, der spielenden Person, ausgelöst; das lokale Dateisystem ist nur beim Sichern, Fortsetzen und im Archiv beteiligt. {tab:uc} beschreibt die Use Cases kurz und ordnet ihnen eine Priorität zu.'),
    ('fig', 'uc', 'abb-use-cases.png', 'Use-Case-Diagramm von ABYSS (UML); hervorgehoben ist der Kernfall UC-02', 13),
    ('tab', 'uc', 'Use Cases mit Kurzbeschreibung und Priorität',
     ['ID', 'Use Case', 'Kurzbeschreibung', 'Priorität'],
     [['UC-01', 'Tauchgang vorbereiten', 'Figur, Waffe und Startbedingungen wählen, Tauchgang beginnen', 'Muss'],
      ['UC-02', 'Zum nächsten Raum vordringen', 'Raum betreten, Gegner besiegen, Raum sichern, Schott öffnen (Kernfall)', 'Muss'],
      ['UC-03', 'Bergung wählen', 'eines von mehreren angebotenen Modulen oder Waffen übernehmen', 'Muss'],
      ['UC-04', 'Handeln', 'Schrott bei Händlerin, Werkstatt oder Kapelle einsetzen', 'Muss'],
      ['UC-05', 'Unterbrechen und fortsetzen', 'pausieren, beenden und am Raumeingang weiterspielen', 'Muss'],
      ['UC-06', 'Brücke erobern', 'letzten Wächter besiegen und einen weiteren Zyklus beginnen', 'Muss'],
      ['UC-07', 'Erneut tauchen', 'nach einer Niederlage die Auswertung ansehen und neu beginnen', 'Muss'],
      ['UC-08', 'Im Archiv freischalten', 'Datenkerne gegen Figuren, Waffen und Module eintauschen', 'Soll'],
      ['UC-09', 'Aussehen anpassen', 'Farben und Helmform der Figur in der Garderobe wählen', 'Soll'],
      ['UC-10', 'Ausrüstung und Karte ansehen', 'installierte Module und den Weg durch das Boot anzeigen', 'Soll'],
      ['UC-11', 'Optionen einstellen', 'Lautstärke, Bildschirmmodus und Effekte anpassen', 'Muss']],
     [.1, .27, .5, .13]),
]

GLOSSARY = [
    ('Bergung', 'Auswahl von Belohnungen nach einem gesicherten Raum; es wird eines von mehreren Angeboten übernommen.'),
    ('Datenkerne', 'Dauerhafte Währung, die über Tauchgänge hinweg erhalten bleibt und im Archiv gegen Freischaltungen eingetauscht wird.'),
    ('Integrität', 'Lebenspunkte der Spielfigur; bei null endet der Tauchgang mit einer Niederlage.'),
    ('Iteration', 'Zeitabschnitt von zwei Wochen, in dem ein überprüfbares Teilergebnis entsteht.'),
    ('Kontextszenario', 'Erzählende Beschreibung, wie eine Person die fertige Anwendung in einer typischen Situation nutzt.'),
    ('Modul', 'Ausrüstungsgegenstand, der Werte oder Fähigkeiten der Spielfigur verändert; aktive Module werden gezielt ausgelöst.'),
    ('Persona', 'Fiktive, aber begründete Beschreibung einer typischen Person der Zielgruppe.'),
    ('Raumtechnik', 'Anlagen in einem Raum, etwa Pressen oder Lasergitter, die Spielfigur und Gegner gleichermassen betreffen.'),
    ('Raumzustand', 'Besondere Regel für einen Raum, etwa Stromausfall oder Hüllenbruch, die vor dem Betreten angekündigt wird.'),
    ('Resonanz', 'Zusatzbonus, der entsteht, wenn zwei bestimmte Module im selben Tauchgang installiert sind.'),
    ('Roguelite', 'Spielgenre mit zufällig zusammengestellten Durchläufen, Neubeginn nach einer Niederlage und teilweise dauerhaftem Fortschritt.'),
    ('Seed', 'Startwert des Zufallsgenerators; derselbe Seed erzeugt dieselbe Abfolge von Räumen.'),
    ('Tauchgang', 'Ein vollständiger Durchlauf vom Heck bis zur Brücke oder bis zur Niederlage.'),
    ('Use Case', 'Anwendungsfall: Beschreibung einer Interaktion zwischen einer Akteurin und dem System mit einem bestimmten Ziel.'),
    ('Wächter', 'Besonders starker Gegner am Ende jeder Sektion.'),
]

KI_TABLE = (
    ['KI', 'Ziel der Verwendung', 'Aufwand für den Prompt', 'Resultat verwendet', 'Art der Verwendung'],
    [['OpenAI Codex', 'Erster spielbarer Prototyp (Version 0.1/0.2), Tests, Figurenmodelle, Klänge', 'Hoch: langer Auftrag mit vielen Prüfschritten', 'Ja', 'Weitgehend übernommen; ab Version 1.0 grösstenteils ersetzt'],
     ['OpenAI Imagegen', 'Konzeptbilder für frühe Präsentationen', 'Mittel', 'Ja', 'Nur in der Konzeptphase, nicht im Spiel'],
     ['Anthropic Claude (Claude Code)', 'Prototyp 1.0/1.1: Spiellogik, Pixelgrafik, Oberfläche, Tests', 'Hoch: autonomer Auftrag mit Prüfschleifen', 'Ja', 'Weitgehend übernommen; wird im Semester vom Team geprüft, erklärt und weiterentwickelt'],
     ['Anthropic Claude (Claude Code)', 'Entwurf dieser Projektskizze, Abbildungen und Präsentationsfolien, Quellenrecherche', 'Mittel', 'Ja', 'Entwurf; Inhalte, Annahmen und Quellen werden vom Team geprüft und verantwortet']],
    [.18, .28, .18, .1, .26])

KI_TEXT = ('Generative KI-Systeme kamen in mehreren Phasen dieses Projekts zum Einsatz {ref:ki_soe}. OpenAI Codex wurde für den ersten spielbaren Prototyp verwendet, OpenAI Imagegen für Konzeptbilder in der Ideenphase. '
           'Anthropic Claude (Claude Code) wurde für die Weiterentwicklung des Prototyps zu Version 1.1 sowie für den Entwurf dieser Projektskizze, ihrer Abbildungen und der Präsentationsfolien eingesetzt. '
           'Die Verantwortung für Inhalt, Richtigkeit und Quellen liegt beim Team. {tab:ki} dokumentiert den Einsatz gemäss den Rahmenbedingungen des Moduls {ref:ki_pm3}.')

# ------------------------------------------------------------------ Nummerierung


class Numbering:
    def __init__(self):
        self.refs: dict[str, int] = {}
        self.figs: dict[str, int] = {}
        self.tabs: dict[str, int] = {}

    def scan(self, text: str):
        for kind, key in re.findall(r'\{(ref|fig|tab):([a-z0-9_]+)\}', text):
            if kind == 'ref' and key not in self.refs:
                self.refs[key] = len(self.refs) + 1

    def build(self):
        for b in CONTENT:
            if b[0] == 'fig':
                self.figs.setdefault(b[1], len(self.figs) + 1)
            if b[0] == 'tab':
                self.tabs.setdefault(b[1], len(self.tabs) + 1)
        self.tabs.setdefault('ki', len(self.tabs) + 1)
        for b in CONTENT:
            for part in b[1:]:
                for t in (part if isinstance(part, list) else [part]):
                    for s in (t if isinstance(t, list) else [t]):
                        if isinstance(s, str):
                            self.scan(s)
        self.scan(KI_TEXT)
        for key in REFS:
            assert key in self.refs, f'Quelle {key} wird nie zitiert'

    def resolve(self, text: str) -> str:
        text = re.sub(r'\{ref:([a-z0-9_]+)\}', lambda m: f'[{self.refs[m.group(1)]}]', text)
        text = re.sub(r'\{fig:([a-z0-9_]+)\}', lambda m: f'Abbildung {self.figs[m.group(1)]}', text)
        return re.sub(r'\{tab:([a-z0-9_]+)\}', lambda m: f'Tabelle {self.tabs[m.group(1)]}', text)


N = Numbering()
N.build()


def headings():
    """Liefert (Nummer, Ebene, Titel) für alle Kapitel inklusive Verzeichnissen."""
    out, h1, h2 = [], 0, 0
    for b in CONTENT:
        if b[0] == 'h1':
            h1, h2 = h1 + 1, 0
            out.append((f'{h1}.', 1, b[1]))
        elif b[0] == 'h2':
            h2 += 1
            out.append((f'{h1}.{h2}.', 2, b[1]))
    tail = ['Quellenverzeichnis', 'Abbildungs- und Tabellenverzeichnisse', 'Glossar', 'Anhang: Deklaration des KI-Einsatzes']
    for t in tail:
        h1 += 1
        out.append((f'{h1}.', 1, t))
        if t.startswith('Abbildungs'):
            out.append((f'{h1}.1.', 2, 'Abbildungsverzeichnis'))
            out.append((f'{h1}.2.', 2, 'Tabellenverzeichnis'))
    return out


def captions():
    figs = [(N.figs[b[1]], b[3]) for b in CONTENT if b[0] == 'fig']
    tabs = [(N.tabs[b[1]], b[2]) for b in CONTENT if b[0] == 'tab'] + [(N.tabs['ki'], 'Einsatz generativer KI im Projekt')]
    return figs, tabs


# ------------------------------------------------------------------ Abbildungen


def pil_font(size, bold=False):
    path = '/System/Library/Fonts/Supplemental/Arial Bold.ttf' if bold else '/System/Library/Fonts/Supplemental/Arial.ttf'
    return ImageFont.truetype(path, size)


def need(path: Path) -> Path:
    if not path.exists():
        raise SystemExit(f'{path} fehlt – zuerst ./gradlew sceneShot --args="build/scenes 0,4,9,14,19 4" ausführen')
    return path


def figure_boot(target: Path):
    src = Image.open(need(SCENES / 'map.png')).convert('RGB')
    S = 2
    top, bottom = 150, 560
    img = src.crop((0, top, 1440, bottom)).resize((1440 * S, (bottom - top) * S), Image.NEAREST)
    strip = 400
    c = Image.new('RGB', (img.width, img.height + strip), 'white')
    c.paste(img)
    d = ImageDraw.Draw(c)
    y0 = img.height
    sections = [('Hecksektion', 180, 413, (201, 128, 40)), ('Maschinendeck', 413, 647, (60, 150, 70)),
                ('Forschungsdeck', 647, 880, (122, 86, 196)), ('Kommandodeck', 880, 1115, (52, 120, 190))]
    for i, (name, a, b, col) in enumerate(sections):
        xa, xb = a * S + 10, b * S - 10
        y = y0 + 40
        d.line([(xa, y), (xb, y)], fill=col, width=9)
        d.line([(xa, y - 26), (xa, y)], fill=col, width=9)
        d.line([(xb, y - 26), (xb, y)], fill=col, width=9)
        cx = (xa + xb) // 2
        d.text((cx, y + 24), name, font=pil_font(52, True), fill=(30, 30, 30), anchor='ma')
        d.text((cx, y + 92), f'Räume {i * 6 + 1}–{i * 6 + 6}', font=pil_font(44), fill=(90, 90, 90), anchor='ma')
    ya = y0 + strip - 64
    d.line([(120, ya), (c.width - 170, ya)], fill=(47, 84, 150), width=10)
    d.polygon([(c.width - 172, ya - 26), (c.width - 120, ya), (c.width - 172, ya + 26)], fill=(47, 84, 150))
    d.text((120, ya - 66), 'Start im Heck', font=pil_font(48, True), fill=(47, 84, 150))
    d.text((c.width - 120, ya - 66), 'Ziel: Brücke', font=pil_font(48, True), fill=(47, 84, 150), anchor='ra')
    c.save(target, optimize=True)


def figure_screen(target: Path):
    im = Image.open(need(SCENES / 'scene-14.png')).convert('RGB')
    S = 2
    im = im.resize((im.width * S, im.height * S), Image.NEAREST)
    d = ImageDraw.Draw(im)
    marks = [(392, 48), (822, 30), (1126, 48), (556, 300), (420, 556), (600, 734), (882, 734)]
    for n, (x, y) in enumerate(marks, 1):
        x, y, r = x * S, y * S, 25 * S
        d.ellipse([x - r - 5, y - r - 5, x + r + 5, y + r + 5], fill=(255, 255, 255))
        d.ellipse([x - r, y - r, x + r, y + r], fill=(47, 84, 150))
        d.text((x, y + 2), str(n), font=pil_font(32 * S, True), fill='white', anchor='mm')
    im.save(target, optimize=True)


def figure_flow(target: Path):
    W, H = 2400, 1000
    c = Image.new('RGB', (W, H), 'white')
    d = ImageDraw.Draw(c)
    blue, grey = (47, 84, 150), (68, 84, 106)

    def node(x, y, w, h, title, sub, fill, edge):
        d.rounded_rectangle([x, y, x + w, y + h], radius=28, fill=fill, outline=edge, width=6)
        d.text((x + w / 2, y + h / 2 - 34), title, font=pil_font(46, True), fill=(25, 25, 25), anchor='mm')
        for i, line in enumerate(sub.split('\n')):
            d.text((x + w / 2, y + h / 2 + 22 + i * 44), line, font=pil_font(36), fill=(70, 70, 70), anchor='mm')
        return (x, y, w, h)

    def arrow(pts, label=None, lpos=None, color=grey, anchor='mm'):
        d.line(pts, fill=color, width=7, joint='curve')
        (x1, y1), (x2, y2) = pts[-2], pts[-1]
        if x1 == x2:
            s = 1 if y2 > y1 else -1
            d.polygon([(x2 - 19, y2 - 34 * s), (x2 + 19, y2 - 34 * s), (x2, y2)], fill=color)
        else:
            s = 1 if x2 > x1 else -1
            d.polygon([(x2 - 34 * s, y2 - 19), (x2 - 34 * s, y2 + 19), (x2, y2)], fill=color)
        if label:
            d.text(lpos, label, font=pil_font(34, True), fill=color, anchor=anchor)

    light, green, red, amber = (234, 240, 250), (231, 244, 234), (251, 234, 234), (255, 244, 224)
    y, h, w = 330, 230, 400
    A = node(30, y, w, h, 'Vorbereiten', 'Figur und Waffe\nwählen', light, blue)
    B = node(510, y, w, h, 'Raum betreten', 'Kampf oder Halt\n(Händler, Werkstatt)', light, blue)
    C = node(990, y, w, h, 'Raum sichern', 'Bergung: Modul\noder Waffe wählen', light, blue)
    D = node(1470, y, w, h, 'Routenwahl', 'einer von zwei\nWegen mit Vorschau', light, blue)
    E = node(1950, y, 420, h, 'Sieg', 'Brücke erobert\nnach Raum 24', green, (60, 141, 79))
    F = node(510, 690, w, 210, 'Niederlage', 'Integrität\nauf null', red, (176, 58, 46))
    G = node(30, 690, w, 210, 'Archiv', 'Datenkerne gegen\nFreischaltungen', amber, (201, 138, 27))
    mid = y + h / 2
    for a, b in ((A, B), (B, C), (C, D), (D, E)):
        arrow([(a[0] + a[2], mid), (b[0], mid)])
    loop_y = 190
    arrow([(D[0] + w / 2, y), (D[0] + w / 2, loop_y), (B[0] + w / 2, loop_y), (B[0] + w / 2, y)],
          'nächster Raum: Schleife über 24 Räume, Wächter in Raum 6, 12, 18 und 24', ((B[0] + D[0] + w) / 2, loop_y - 50), blue)
    arrow([(B[0] + w / 2, y + h), (B[0] + w / 2, F[1])], 'im Kampf', (B[0] + w / 2 + 22, (y + h + F[1]) / 2), (176, 58, 46), 'lm')
    arrow([(F[0], F[1] + 105), (G[0] + w, G[1] + 105)])
    arrow([(E[0] + 210, y + h), (E[0] + 210, 960), (G[0] + w / 2, 960), (G[0] + w / 2, G[1] + 210)],
          'nach dem Sieg: nächster Zyklus möglich', (1560, 922), (60, 141, 79))
    arrow([(G[0] + 90, G[1]), (G[0] + 90, y + h)], 'neuer Tauchgang', (G[0] + 112, (y + h + G[1]) / 2), (201, 138, 27), 'lm')
    c.save(target, optimize=True)


def figure_uc(target: Path):
    """UML-Use-Case-Diagramm nach docs/diagrams/01-use-cases.puml, für den Druck gross beschriftet."""
    W, H = 2300, 1720
    c = Image.new('RGB', (W, H), 'white')
    d = ImageDraw.Draw(c)
    ink, edge, fill = (40, 48, 58), (53, 88, 99), (234, 243, 242)
    cases = [('UC-01', 'Tauchgang vorbereiten'), ('UC-02', 'Zum nächsten Raum vordringen'), ('UC-03', 'Bergung wählen'),
             ('UC-04', 'Handeln'), ('UC-05', 'Unterbrechen und fortsetzen'), ('UC-06', 'Brücke erobern'), ('UC-07', 'Erneut tauchen'),
             ('UC-08', 'Im Archiv freischalten'), ('UC-09', 'Aussehen anpassen'), ('UC-10', 'Ausrüstung und Karte ansehen'),
             ('UC-11', 'Optionen einstellen')]
    bx0, by0, bx1, by1 = 470, 70, 1640, H - 40
    d.rectangle([bx0, by0, bx1, by1], outline=ink, width=4)
    d.text((bx0 + 24, by0 + 18), 'ABYSS 1.1 · lokales Desktopspiel', font=pil_font(40, True), fill=ink)
    cx, ow, oh, gap = 1000, 780, 112, 138
    centers = [(cx, 210 + i * gap) for i in range(len(cases))]

    def actor(x, y, label):
        d.ellipse([x - 34, y - 190, x + 34, y - 122], outline=ink, width=6)
        d.line([(x, y - 122), (x, y - 20)], fill=ink, width=6)
        d.line([(x - 70, y - 88), (x + 70, y - 88)], fill=ink, width=6)
        d.line([(x, y - 20), (x - 55, y + 60)], fill=ink, width=6)
        d.line([(x, y - 20), (x + 55, y + 60)], fill=ink, width=6)
        for i, part in enumerate(label.split('\n')):
            d.text((x, y + 80 + i * 48), part, font=pil_font(42, True), fill=ink, anchor='ma')

    ax, ay = 210, H // 2
    for x, y in centers:
        d.line([(ax + 60, ay - 90), (x - ow / 2, y)], fill=edge, width=4)
    fx, fy = 2010, 820
    for i in (0, 4, 7):
        x, y = centers[i]
        d.line([(x + ow / 2, y), (fx - 70, fy - 90)], fill=edge, width=4)
    for i in (2, 3):
        x, y = centers[i]
        tx, ty = centers[1]
        bulge = 1500 + (i - 2) * 70
        pts = [(x + ow / 2 - 30, y - oh / 2 + 12)]
        steps = 24
        for k in range(1, steps + 1):
            t = k / steps
            px = (1 - t) ** 2 * (x + ow / 2 - 30) + 2 * (1 - t) * t * bulge + t ** 2 * (tx + ow / 2 - 30)
            py = (1 - t) ** 2 * (y - oh / 2 + 12) + 2 * (1 - t) * t * ((y + ty) / 2) + t ** 2 * (ty + oh / 2 - 12)
            pts.append((px, py))
        for k in range(0, len(pts) - 1, 2):
            d.line([pts[k], pts[k + 1]], fill=edge, width=4)
        ex, ey = pts[-1]
        px, py = pts[-3]
        import math
        ang = math.atan2(ey - py, ex - px)
        for sgn in (-1, 1):
            d.line([(ex, ey), (ex - 30 * math.cos(ang + sgn * .45), ey - 30 * math.sin(ang + sgn * .45))], fill=edge, width=4)
        d.text((bulge + 34, (y + ty) / 2 + (-8 if i == 2 else 22)), '«extend»', font=pil_font(34), fill=edge, anchor='lm')
    for (x, y), (uid, name) in zip(centers, cases):
        core = uid == 'UC-02'
        d.ellipse([x - ow / 2, y - oh / 2, x + ow / 2, y + oh / 2], fill=(255, 241, 214) if core else fill,
                  outline=(196, 128, 30) if core else edge, width=5)
        d.text((x, y - 24), uid, font=pil_font(32), fill=(90, 96, 100), anchor='mm')
        d.text((x, y + 18), name, font=pil_font(42, True), fill=ink, anchor='mm')
    actor(ax, ay, 'Spieler/in')
    actor(fx, fy, 'Lokales\nDateisystem')
    c.save(target, optimize=True)


def build_figures():
    FIG.mkdir(parents=True, exist_ok=True)
    figure_boot(FIG / 'abb-boot.png')
    figure_screen(FIG / 'abb-bildschirm.png')
    figure_flow(FIG / 'abb-ablauf.png')
    figure_uc(FIG / 'abb-use-cases.png')
    title = Image.open(need(SCENES / 'title.png')).convert('RGB')
    title.save(FIG / 'titelbild.png', optimize=True)


# ------------------------------------------------------------------ HTML / PDF


def inline_html(text: str) -> str:
    text = N.resolve(text)
    parts = re.split(r'(<i>.*?</i>)', text)
    out = []
    for p in parts:
        if p.startswith('<i>'):
            out.append(p)
            continue
        p = html.escape(p, quote=False)
        p = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', p)
        p = re.sub(r'\*(.+?)\*', r'<em>\1</em>', p)
        out.append(p)
    return ''.join(out)


def font_face():
    faces = []
    for name, file, weight, style in [('Calibri Light', 'calibril.ttf', 300, 'normal'), ('Calibri', 'Calibri.ttf', 400, 'normal'),
                                      ('Calibri', 'Calibrib.ttf', 700, 'normal')]:
        p = WORD_FONTS / file
        if p.exists():
            faces.append(f"@font-face {{ font-family: '{name}'; src: url('{p.as_uri()}'); font-weight: {weight}; font-style: {style}; }}")
    return '\n'.join(faces)


CSS = """
@page { size: A4; margin: 25mm 25mm 27mm 25mm;
  @bottom-left { content: "Bachelor of Science (BSc) in Informatik\\A Software-Projekt 3 (PM3)\\A Projektskizze %TEAM%"; white-space: pre; font: 8pt Arial, sans-serif; color: #404040; vertical-align: top; padding-top: 5mm; }
  @bottom-right { content: counter(page) "/" counter(pages); font: 8pt Arial, sans-serif; color: #404040; vertical-align: top; padding-top: 12.2mm; }
}
@page :first { @bottom-left { content: none; } @bottom-right { content: none; } }
html { font-family: Arial, Helvetica, sans-serif; font-size: 11pt; color: #1a1a1a; }
body { margin: 0; line-height: 1.38; hyphens: auto; -webkit-hyphens: auto; }
p { margin: 0 0 7pt 0; text-align: justify; orphans: 3; widows: 3; }
h1, h2, .tochead { font-family: 'Calibri Light', 'Calibri', Arial, sans-serif; font-weight: 300; color: #2F5496; }
h1 { font-size: 14pt; margin: 18pt 0 9pt 0; break-after: avoid; }
h2 { font-size: 12pt; margin: 12pt 0 6pt 0; break-after: avoid; }
h1 .num, h2 .num { display: inline-block; min-width: 12mm; }
.title { height: 245mm; display: flex; flex-direction: column; align-items: center; text-align: center; break-after: page; }
.title .l14 { font-size: 14pt; font-weight: bold; margin: 0 0 8pt 0; }
.title .l18 { font-size: 18pt; font-weight: bold; margin: 34pt 0 14pt 0; }
.title .l20 { font-size: 20pt; font-weight: bold; margin: 30pt 0 6pt 0; }
.title .sub { font-size: 13pt; color: #404040; margin: 0 0 20pt 0; }
.title img { width: 128mm; border: 0.6pt solid #1a1a1a; }
.title .date { margin-top: 22pt; }
.title .names { margin-top: 16pt; line-height: 1.7; }
.title .todo { color: #8a8a8a; font-style: italic; font-size: 10pt; }
.toc { break-after: page; }
.tochead { font-size: 16pt; margin: 0 0 14pt 0; }
.toc .e { display: flex; align-items: baseline; margin: 0 0 5pt 0; }
.toc .e.l2 { padding-left: 8mm; }
.toc .e .n { min-width: 12mm; }
.toc .e .dots { flex: 1; border-bottom: 0.8pt dotted #7a7a7a; margin: 0 4pt; transform: translateY(-3pt); }
figure { margin: 10pt 0 12pt 0; text-align: center; break-inside: avoid; }
figure img { max-width: 100%; }
figcaption, .tcap { font-size: 9pt; color: #44546A; font-style: italic; text-align: left; }
figcaption { margin-top: 5pt; }
.tcap { margin: 10pt 0 4pt 0; break-after: avoid; }
table { width: 100%; border-collapse: collapse; font-size: 9pt; line-height: 1.3; margin: 0 0 12pt 0; break-inside: auto; }
thead { display: table-header-group; }
table.keep { break-inside: avoid; }
tr { break-inside: avoid; }
th, td { border: 0.6pt solid #808080; padding: 3.5pt 5pt; vertical-align: top; text-align: left; hyphens: auto; }
th { background: #D9E2F3; font-weight: bold; }
ul { margin: 0 0 8pt 0; padding-left: 6mm; }
li { margin: 0 0 4pt 0; text-align: justify; }
.persona { border-left: 3pt solid #2F5496; background: #F2F5FB; padding: 7pt 10pt; margin: 6pt 0 10pt 0; }
.persona p { margin: 0; }
.refs p { text-align: left; padding-left: 9mm; text-indent: -9mm; font-size: 10pt; margin-bottom: 5pt; word-break: break-word; }
.refs .rn { display: inline-block; width: 9mm; text-indent: 0; }
.lof .e { display: flex; align-items: baseline; margin: 0 0 4pt 0; font-size: 10pt; }
.lof .e .dots { flex: 1; border-bottom: 0.8pt dotted #7a7a7a; margin: 0 4pt; transform: translateY(-3pt); }
.gl p { text-align: left; margin-bottom: 4pt; }
"""


def build_html(pages: dict[str, int] | None) -> str:
    pg = pages or {}
    H = []
    H.append('<!doctype html><html lang="de-CH"><head><meta charset="utf-8"><title>Projektskizze ABYSS</title><style>')
    H.append(font_face())
    H.append(CSS.replace('%TEAM%', TEAM))
    H.append('</style></head><body>')
    # Titelblatt
    names = ''.join(f'<div>{html.escape(m)}</div>' for m in MEMBERS)
    todo = f'<div class="todo">[{html.escape(MEMBERS_TODO)}]</div>' if MEMBERS_TODO else ''
    H.append(f'''<section class="title">
      <div class="l14" style="margin-top:6mm">ZHAW School of Engineering</div>
      <div class="l14">Bachelor of Science in Informatik</div>
      <div class="l14">Software-Projekt PM3 HS26</div>
      <div class="l18">Projektskizze</div>
      <div class="l14">{html.escape(TEAM)}</div>
      <div class="l20">{html.escape(PRODUCT)}</div>
      <div class="sub">{html.escape(SUBTITLE)}</div>
      <img src="{(FIG / 'titelbild.png').as_uri()}" alt="Titelbild">
      <div class="date">Abgabedatum: {DATE}</div>
      <div class="names">Namen der Team-Mitglieder in alphabetischer Reihenfolge:{names}{todo}</div>
    </section>''')
    # Inhaltsverzeichnis
    H.append('<section class="toc"><div class="tochead">Inhaltsverzeichnis</div>')
    for num, level, title in headings():
        H.append(f'<div class="e l{level}"><span class="n">{num}</span><span>{html.escape(title)}</span><span class="dots"></span><span>{pg.get("h:" + num, 0)}</span></div>')
    H.append('</section>')
    # Kapitel
    hs = iter(headings())
    for b in CONTENT:
        kind = b[0]
        if kind in ('h1', 'h2'):
            num, level, title = next(hs)
            H.append(f'<h{level}><span class="num">{num}</span>{html.escape(title)}</h{level}>')
        elif kind == 'p':
            H.append(f'<p>{inline_html(b[1])}</p>')
        elif kind == 'persona':
            H.append(f'<div class="persona"><p>{inline_html(b[1])}</p></div>')
        elif kind == 'ul':
            H.append('<ul>' + ''.join(f'<li>{inline_html(x)}</li>' for x in b[1]) + '</ul>')
        elif kind == 'fig':
            _, key, file, cap, width = b
            H.append(f'<figure><img src="{(FIG / file).as_uri()}" style="width:{width}cm"><figcaption>Abbildung {N.figs[key]}: {inline_html(cap)}</figcaption></figure>')
        elif kind == 'tab':
            _, key, cap, head, rows, widths = b
            H.append(table_html(N.tabs[key], cap, head, rows, widths))
    rest = list(hs)
    # Quellen
    H.append(f'<h1><span class="num">{rest[0][0]}</span>{rest[0][2]}</h1><div class="refs">')
    for key, n in sorted(N.refs.items(), key=lambda kv: kv[1]):
        H.append(f'<p><span class="rn">[{n}]</span>{REFS[key]}</p>')
    H.append('</div>')
    # Verzeichnisse
    figs, tabs = captions()
    H.append(f'<h1><span class="num">{rest[1][0]}</span>{rest[1][2]}</h1>')
    H.append(f'<h2><span class="num">{rest[2][0]}</span>{rest[2][2]}</h2><div class="lof">')
    for n, cap in figs:
        H.append(f'<div class="e"><span>Abbildung {n}: {inline_html(cap)}</span><span class="dots"></span><span>{pg.get(f"f:{n}", 0)}</span></div>')
    H.append(f'</div><h2><span class="num">{rest[3][0]}</span>{rest[3][2]}</h2><div class="lof">')
    for n, cap in tabs:
        H.append(f'<div class="e"><span>Tabelle {n}: {inline_html(cap)}</span><span class="dots"></span><span>{pg.get(f"t:{n}", 0)}</span></div>')
    H.append('</div>')
    # Glossar
    H.append(f'<h1><span class="num">{rest[4][0]}</span>{rest[4][2]}</h1><div class="gl">')
    for term, text in GLOSSARY:
        H.append(f'<p><strong>{html.escape(term)}:</strong> {html.escape(text)}</p>')
    H.append('</div>')
    # KI
    H.append(f'<h1><span class="num">{rest[5][0]}</span>{rest[5][2]}</h1><p>{inline_html(KI_TEXT)}</p>')
    H.append(table_html(N.tabs['ki'], 'Einsatz generativer KI im Projekt (Darstellung nach [%d])' % N.refs['ki_pm3'], *KI_TABLE))
    H.append('</body></html>')
    return '\n'.join(H)


def table_html(n, cap, head, rows, widths):
    keep = ' class="keep"' if len(rows) <= 8 else ''
    t = [f'<div class="tcap">Tabelle {n}: {inline_html(cap)}</div><table{keep}><colgroup>']
    t += [f'<col style="width:{w * 100:.1f}%">' for w in widths]
    t.append('</colgroup><thead><tr>' + ''.join(f'<th>{inline_html(h)}</th>' for h in head) + '</tr></thead><tbody>')
    for r in rows:
        t.append('<tr>' + ''.join(f'<td>{inline_html(c)}</td>' for c in r) + '</tr>')
    t.append('</tbody></table>')
    return ''.join(t)


def chrome_pdf(html_path: Path, pdf_path: Path):
    """Druckt die HTML-Seite mit Chrome als PDF. Chrome bleibt nach dem Druck teils offen, daher wird es beendet."""
    import time
    profile = BUILD / 'chrome-profile'
    if pdf_path.exists():
        pdf_path.unlink()
    proc = subprocess.Popen([CHROME, '--headless', '--disable-gpu', '--no-pdf-header-footer', '--allow-file-access-from-files',
                             '--no-first-run', '--no-default-browser-check', f'--user-data-dir={profile}',
                             f'--print-to-pdf={pdf_path}', html_path.as_uri()],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    deadline, last = time.time() + 90, -1
    try:
        while time.time() < deadline:
            if proc.poll() is not None and pdf_path.exists():
                break
            if pdf_path.exists():
                size = pdf_path.stat().st_size
                if size == last and size > 0 and pdf_path.read_bytes()[-1024:].rstrip().endswith(b'%%EOF'):
                    break
                last = size
            time.sleep(.5)
        else:
            raise SystemExit('Chrome hat kein PDF erzeugt')
    finally:
        if proc.poll() is None:
            proc.terminate()
            try:
                proc.wait(5)
            except subprocess.TimeoutExpired:
                proc.kill()


def page_map(pdf: Path) -> dict[str, int]:
    count = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', str(pdf)], capture_output=True, text=True).stdout).group(1))
    found: dict[str, int] = {}
    figs, tabs = captions()
    for page in range(3, count + 1):
        text = subprocess.run(['pdftotext', '-f', str(page), '-l', str(page), '-layout', str(pdf), '-'], capture_output=True, text=True).stdout
        lines = [ln.strip() for ln in text.splitlines()]
        for num, level, title in headings():
            key = 'h:' + num
            if key not in found and any(re.match(re.escape(num) + r'\s+' + re.escape(title[:18]), ln) for ln in lines):
                found[key] = page
        for n, _ in figs:
            if f'f:{n}' not in found and any(ln.startswith(f'Abbildung {n}:') for ln in lines):
                found[f'f:{n}'] = page
        for n, _ in tabs:
            if f't:{n}' not in found and any(ln.startswith(f'Tabelle {n}:') for ln in lines):
                found[f't:{n}'] = page
    missing = [k for k in [f'h:{h[0]}' for h in headings()] if k not in found]
    if missing:
        print('Seitenzahl nicht gefunden für', missing)
    return found


# ------------------------------------------------------------------ Word (SoE-Vorlage)


def build_docx(pages: dict[str, int], target: Path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = docx.Document(str(TEMPLATE))
    P = doc.paragraphs

    def set_text(p, text):
        runs = p.runs
        runs[0].text = text
        for r in runs[1:]:
            r.text = ''

    set_text(P[2], 'Software-Projekt PM3 HS26')
    set_text(P[5], TEAM)
    set_text(P[8], PRODUCT)
    sub = P[9]
    r = sub.add_run(SUBTITLE)
    r.font.size = Pt(13)
    set_text(P[10], f'Abgabedatum: {DATE}')
    names_anchor = P[12]
    for m in MEMBERS:
        np_ = names_anchor.insert_paragraph_before(m)
        np_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if MEMBERS_TODO:
        np_ = names_anchor.insert_paragraph_before()
        np_.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = np_.add_run(f'[{MEMBERS_TODO}]')
        rr.italic = True
        rr.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
    # Titelbild zwischen Produkttitel und Datum
    pic_p = P[10].insert_paragraph_before()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(FIG / 'titelbild.png'), width=Cm(12.5))

    body = doc.element.body
    sdt = next(el for el in body.iterchildren() if el.tag == qn('w:sdt')
               and any(g.get(qn('w:val')) == 'Table of Contents' for g in el.iter(qn('w:docPartGallery'))))
    # alles nach dem Inhaltsverzeichnis entfernen (ausser Abschnittseigenschaften)
    after = False
    for el in list(body.iterchildren()):
        if el is sdt:
            after = True
            continue
        if after and el.tag != qn('w:sectPr'):
            body.remove(el)

    def fld(kind):
        r = OxmlElement('w:r')
        f = OxmlElement('w:fldChar')
        f.set(qn('w:fldCharType'), kind)
        r.append(f)
        return r

    def instr(text):
        r = OxmlElement('w:r')
        t = OxmlElement('w:instrText')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        r.append(t)
        return r

    def text_run(text, bold=False, italic=False, tab=False):
        r = OxmlElement('w:r')
        if bold or italic:
            rpr = OxmlElement('w:rPr')
            if bold:
                rpr.append(OxmlElement('w:b'))
            if italic:
                rpr.append(OxmlElement('w:i'))
            r.append(rpr)
        if tab:
            r.append(OxmlElement('w:tab'))
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        r.append(t)
        return r

    def field_list(container_append, entries, instruction, style_for):
        """Baut ein Verzeichnisfeld mit zwischengespeicherten Einträgen; Word aktualisiert es beim Öffnen."""
        for i, (level, left, title, page) in enumerate(entries):
            p = OxmlElement('w:p')
            ppr = OxmlElement('w:pPr')
            ps = OxmlElement('w:pStyle')
            ps.set(qn('w:val'), style_for(level))
            ppr.append(ps)
            tabs = OxmlElement('w:tabs')
            tb = OxmlElement('w:tab')
            tb.set(qn('w:val'), 'right')
            tb.set(qn('w:leader'), 'dot')
            tb.set(qn('w:pos'), '9060')
            tabs.append(tb)
            ppr.append(tabs)
            p.append(ppr)
            if i == 0:
                p.append(fld('begin'))
                p.append(instr(instruction))
                p.append(fld('separate'))
            if left:
                p.append(text_run(left))
                p.append(text_run(title, tab=True))
            else:
                p.append(text_run(title))
            p.append(text_run(str(page or ''), tab=True))
            if i == len(entries) - 1:
                p.append(fld('end'))
            container_append(p)

    # Inhaltsverzeichnis neu füllen
    content = sdt.find(qn('w:sdtContent'))
    paras = content.findall(qn('w:p'))
    for p in paras[1:]:
        content.remove(p)
    style_id = {s.name: s.style_id for s in doc.styles}
    field_list(content.append, [(lvl, num, title, pages.get('h:' + num)) for num, lvl, title in headings()],
               ' TOC \\o "1-3" \\h \\z \\u ', lambda lvl: style_id.get(f'toc {lvl}', 'Verzeichnis1'))

    def add_inline(p, text):
        text = N.resolve(text)
        text = re.sub(r'</?i>', '*', text)
        for part in re.split(r'(\*\*.+?\*\*|\*.+?\*)', text):
            if not part:
                continue
            if part.startswith('**'):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith('*'):
                p.add_run(part[1:-1]).italic = True
            else:
                p.add_run(part)

    def caption(kind, n, text):
        p = doc.add_paragraph(style='Caption')
        p.add_run(f'{kind} ')
        fs = OxmlElement('w:fldSimple')
        fs.set(qn('w:instr'), f' SEQ {kind} \\* ARABIC ')
        fs.append(text_run(str(n)))
        p._p.append(fs)
        add_inline(p, ': ' + text)
        if kind == 'Tabelle':
            p.paragraph_format.keep_with_next = True
        return p

    def shade(cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill)
        tcpr.append(shd)

    def table(n, cap, head, rows, widths):
        caption('Tabelle', n, cap)
        t = doc.add_table(rows=1 + len(rows), cols=len(head))
        t.style = doc.styles['Table Grid']
        total = 16.0
        for ri, row in enumerate([head] + rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.width = Cm(total * widths[ci])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                add_inline(p, f'**{val}**' if ri == 0 and not val.startswith('**') else val)
                for run in p.runs:
                    run.font.size = Pt(9)
                if ri == 0:
                    shade(cell, 'D9E2F3')
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader')
        trpr.append(th)
        doc.add_paragraph(style='Standard Technischer Bericht')

    for b in CONTENT:
        kind = b[0]
        if kind == 'h1':
            doc.add_paragraph(b[1], style='Heading 1')
        elif kind == 'h2':
            doc.add_paragraph(b[1], style='Heading 2')
        elif kind in ('p', 'persona'):
            p = doc.add_paragraph(style='Standard Technischer Bericht')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, b[1])
            if kind == 'persona':
                p.paragraph_format.left_indent = Cm(0.6)
                ppr = p._p.get_or_add_pPr()
                bdr = OxmlElement('w:pBdr')
                left = OxmlElement('w:left')
                for k, v in (('w:val', 'single'), ('w:sz', '18'), ('w:space', '8'), ('w:color', '2F5496')):
                    left.set(qn(k), v)
                bdr.append(left)
                ppr.append(bdr)
        elif kind == 'ul':
            for item in b[1]:
                p = doc.add_paragraph(style='Standard Technischer Bericht')
                p.paragraph_format.left_indent = Cm(0.63)
                p.paragraph_format.first_line_indent = Cm(-0.63)
                p.paragraph_format.space_after = Pt(3)
                p.add_run('•\t')
                add_inline(p, item)
        elif kind == 'fig':
            _, key, file, cap, width = b
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(FIG / file), width=Cm(width))
            caption('Abbildung', N.figs[key], cap)
        elif kind == 'tab':
            _, key, cap, head, rows, widths = b
            table(N.tabs[key], cap, head, rows, widths)

    tail = [h for h in headings() if h[2] in ('Quellenverzeichnis', 'Abbildungs- und Tabellenverzeichnisse', 'Abbildungsverzeichnis',
                                              'Tabellenverzeichnis', 'Glossar', 'Anhang: Deklaration des KI-Einsatzes')]
    doc.add_paragraph('Quellenverzeichnis', style='Heading 1')
    for key, n in sorted(N.refs.items(), key=lambda kv: kv[1]):
        p = doc.add_paragraph(style='Bibliographie')
        p.add_run(f'[{n}]\t')
        add_inline(p, REFS[key])
    figs, tabs = captions()
    doc.add_paragraph('Abbildungs- und Tabellenverzeichnisse', style='Heading 1')
    doc.add_paragraph('Abbildungsverzeichnis', style='Heading 2')
    tof = style_id.get('table of figures', 'Abbildungsverzeichnis')
    body_append = lambda el: body.insert(len(body) - 1, el)
    field_list(body_append, [(1, None, f'Abbildung {n}: {N.resolve(c)}', pages.get(f'f:{n}')) for n, c in figs], ' TOC \\h \\z \\c "Abbildung" ', lambda _: tof)
    doc.add_paragraph('Tabellenverzeichnis', style='Heading 2')
    field_list(body_append, [(1, None, f'Tabelle {n}: {N.resolve(c)}', pages.get(f't:{n}')) for n, c in tabs], ' TOC \\h \\z \\c "Tabelle" ', lambda _: tof)
    doc.add_paragraph('Glossar', style='Heading 1')
    for term, text in GLOSSARY:
        p = doc.add_paragraph(style='Standard Technischer Bericht')
        p.paragraph_format.space_after = Pt(4)
        p.add_run(term + ': ').bold = True
        p.add_run(text)
    doc.add_paragraph('Anhang: Deklaration des KI-Einsatzes', style='Heading 1')
    p = doc.add_paragraph(style='Standard Technischer Bericht')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, KI_TEXT)
    table(N.tabs['ki'], 'Einsatz generativer KI im Projekt (Darstellung nach [%d])' % N.refs['ki_pm3'], *KI_TABLE)

    # Fusszeile
    for sec in doc.sections:
        for fp in sec.footer.paragraphs:
            for run in fp.runs:
                if 'xxx' in run.text:
                    run.text = run.text.replace('Team xxx', TEAM).replace('xxx', TEAM.replace('Team ', ''))
    # Felder beim Öffnen aktualisieren
    settings = doc.settings.element
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    settings.append(uf)
    doc.core_properties.title = f'Projektskizze {PRODUCT}'
    doc.core_properties.subject = 'PM3 HS26 · Meilenstein M1'
    doc.core_properties.author = ', '.join(MEMBERS)
    doc.save(str(target))


# ------------------------------------------------------------------ Markdown (GitHub)


def md_inline(text):
    text = N.resolve(text)
    return re.sub(r'</?i>', '*', text)


def build_md(target: Path):
    L = [f'# Projektskizze {PRODUCT}', '',
         f'*PM3 HS26 · Meilenstein M1 · {TEAM} · Abgabe {DATE}. Lesefassung; massgeblich ist [{STEM}.pdf]({STEM}.pdf) '
         f'(Word: [{STEM}.docx]({STEM}.docx)). Erzeugt mit `tools/create_projektskizze.py`.*', '']
    hs = iter(headings())
    for b in CONTENT:
        kind = b[0]
        if kind in ('h1', 'h2'):
            num, level, title = next(hs)
            L += [f'{"#" * (level + 1)} {num} {title}', '']
        elif kind == 'p':
            L += [md_inline(b[1]), '']
        elif kind == 'persona':
            L += ['> ' + md_inline(b[1]), '']
        elif kind == 'ul':
            L += [f'- {md_inline(x)}' for x in b[1]] + ['']
        elif kind == 'fig':
            _, key, file, cap, _w = b
            L += [f'![Abbildung {N.figs[key]}](abbildungen/{file})', '', f'*Abbildung {N.figs[key]}: {md_inline(cap)}*', '']
        elif kind == 'tab':
            _, key, cap, head, rows, _w = b
            L += [f'*Tabelle {N.tabs[key]}: {md_inline(cap)}*', '', '| ' + ' | '.join(head) + ' |', '|' + '---|' * len(head)]
            L += ['| ' + ' | '.join(md_inline(c) for c in r) + ' |' for r in rows] + ['']
    rest = list(hs)
    L += [f'## {rest[0][0]} Quellenverzeichnis', '']
    for key, n in sorted(N.refs.items(), key=lambda kv: kv[1]):
        L += [f'[{n}] {md_inline(REFS[key])}', '']
    L += [f'## {rest[4][0]} Glossar', ''] + [f'- **{t}:** {x}' for t, x in GLOSSARY] + ['']
    L += [f'## {rest[5][0]} Anhang: Deklaration des KI-Einsatzes', '', md_inline(KI_TEXT), '']
    head, rows, _w = KI_TABLE
    L += [f'*Tabelle {N.tabs["ki"]}: Einsatz generativer KI im Projekt*', '', '| ' + ' | '.join(head) + ' |', '|' + '---|' * len(head)]
    L += ['| ' + ' | '.join(r) + ' |' for r in rows] + ['']
    target.write_text('\n'.join(L), encoding='utf-8')


# ------------------------------------------------------------------ Ablauf


def main():
    BUILD.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    build_figures()
    html_path = BUILD / 'projektskizze.html'
    html_path.write_text(build_html(None), encoding='utf-8')
    first = BUILD / 'pass1.pdf'
    chrome_pdf(html_path, first)
    pages = page_map(first)
    html_path.write_text(build_html(pages), encoding='utf-8')
    pdf = OUT / f'{STEM}.pdf'
    chrome_pdf(html_path, pdf)
    check = page_map(pdf)
    if check != pages:
        print('Hinweis: Seitenzahlen haben sich im zweiten Lauf verschoben', {k: (pages.get(k), v) for k, v in check.items() if pages.get(k) != v})
    build_docx(pages, OUT / f'{STEM}.docx')
    build_md(OUT / 'PROJEKTSKIZZE.md')
    count = subprocess.run(['pdfinfo', str(pdf)], capture_output=True, text=True).stdout
    print(pdf.relative_to(ROOT), re.search(r'Pages:\s+(\d+)', count).group(1), 'Seiten')
    print((OUT / f'{STEM}.docx').relative_to(ROOT))


if __name__ == '__main__':
    main()
